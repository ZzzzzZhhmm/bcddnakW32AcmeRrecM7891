"""Build the partner Word handoff from its Markdown source (authoring only).

Requires python-docx. Render and inspect every page before distributing changes.
No runtime/robot environment needs this optional document-authoring dependency.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/real/PARTNER_HANDOFF_ZH.md"
OUTPUT = SOURCE.with_suffix(".docx")


def font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def inline(paragraph, text, size=11):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = part.startswith("**")
        clean = part[2:-2] if bold else part.strip("`")
        font(paragraph.add_run(clean), size, bold)


def set_table(table, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    count = len(rows[0])
    widths = ([3.25, 3.85] if count == 2 else [1.4, 2.65, 3.05] if count == 3 else [1.25, 2.0, 1.9, 1.95])
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:"+edge)
        for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn("w:"+key), val)
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for i, row_values in enumerate(rows):
        row = table.rows[i]
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if i == 0:
            props.append(OxmlElement("w:tblHeader"))
        for j, text in enumerate(row_values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge, val in (("top", "70"), ("bottom", "70"), ("left", "90"), ("right", "90")):
                element = OxmlElement("w:"+edge)
                element.set(qn("w:w"), val)
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tcpr.append(margins)
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DCE6F1")
                tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(14.5)
            if i == 0:
                p.paragraph_format.keep_with_next = True
            inline(p, text, 10.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
    section.left_margin, section.right_margin = Inches(.7), Inches(.7)
    section.header_distance = Inches(.25)
    section.footer_distance = Inches(.25)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = Pt(16)
    normal.paragraph_format.space_after = Pt(6)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    normal._element.get_or_add_pPr().append(snap)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = Pt(size+6)
    # The bundled template can carry Word title borders; remove decoration.
    for element in doc.styles.element.xpath(".//w:pBdr"):
        element.getparent().remove(element)
    header = section.header.paragraphs[0]
    font(header.add_run("WARM  真机协作说明"), 8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    font(footer.add_run("WARM  /  "), 8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "WARM Piper 真机部署与实验协作说明"
    doc.core_properties.author = "WARM 项目组"
    doc.core_properties.subject = "第三方现场条件确认 数据采集 推理部署与实验协作"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    page_before_next = False
    while i < len(lines):
        line = lines[i]
        i += 1
        if not line.strip():
            continue
        if line == "<!-- pagebreak -->":
            page_before_next = True
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], "Heading 1")
            if page_before_next:
                p.paragraph_format.page_break_before = True
                page_before_next = False
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("|"):
            rows = [[x.strip() for x in line.strip("|").split("|")]]
            while i < len(lines) and lines[i].startswith("|"):
                raw = lines[i]
                i += 1
                if re.fullmatch(r"[| :\-]+", raw):
                    continue
                rows.append([x.strip() for x in raw.strip("|").split("|")])
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            set_table(table, rows)
            spacer = doc.add_paragraph()
            spacer.paragraph_format.line_spacing = Pt(2)
            spacer.paragraph_format.space_after = Pt(3)
        elif line.startswith("```"):
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                run = p.add_run(lines[i])
                font(run, 9)
                run.font.name = "Consolas"
                i += 1
            i += 1
        else:
            inline(doc.add_paragraph(), line)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
